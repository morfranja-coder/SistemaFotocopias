from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def export_docx(text: str, output_path: str | Path) -> Path:
    output = Path(output_path)
    document = Document()
    for block in text.split("\n\n"):
        paragraph = document.add_paragraph()
        lines = block.splitlines() or [""]
        for index, line in enumerate(lines):
            if index:
                paragraph.add_run().add_break()
            paragraph.add_run(line)
    document.save(output)
    return output


def export_pdf(text: str, output_path: str | Path) -> Path:
    output = Path(output_path)
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    body.leading = 15

    story = []
    blocks = text.split("\n\n") if text else [""]
    for block in blocks:
        safe = (
            block.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace("\n", "<br/>")
        )
        story.append(Paragraph(safe or " ", body))
        story.append(Spacer(1, 8))

    doc = SimpleDocTemplate(
        str(output),
        pagesize=A4,
        rightMargin=50,
        leftMargin=50,
        topMargin=50,
        bottomMargin=50,
        title=output.stem,
    )
    doc.build(story)
    return output
